"""ساخت فایل‌های Excel واقعی (XLSX) برای خروجی‌های پنل ادمین."""

from __future__ import annotations

import csv
import io
import mimetypes
from collections.abc import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
mimetypes.add_type(XLSX_MIME_TYPE, ".xlsx")

_HEADER_FILL = PatternFill("solid", fgColor="0F464F")
_HEADER_FONT = Font(name="Aptos", color="FFFFFF", bold=True)
_EVEN_FILL = PatternFill("solid", fgColor="F2F8F6")
_BOTTOM_BORDER = Border(bottom=Side(style="hair", color="CCDDD9"))


def xlsx_bytes(
    headers: Sequence[object],
    rows: Iterable[Sequence[object]],
    *,
    sheet_name: str = "خروجی",
) -> bytes:
    """تولید workbook تک‌برگی راست‌به‌چپ؛ متن‌ها هرگز به‌عنوان فرمول اجرا نمی‌شوند."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_name.strip() or "خروجی")[:31]
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    all_rows = [list(headers), *(list(row) for row in rows)]
    for row_index, values in enumerate(all_rows, start=1):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = value
            if isinstance(value, str):
                # جلوگیری از تفسیر متن ثبت‌شده توسط کاربر به‌عنوان فرمول اکسل.
                cell.data_type = "s"
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if row_index == 1:
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.border = _BOTTOM_BORDER
                if row_index % 2 == 0:
                    cell.fill = _EVEN_FILL

    sheet.row_dimensions[1].height = 25
    if all_rows and all_rows[0]:
        sheet.auto_filter.ref = sheet.dimensions
        for column_index in range(1, len(all_rows[0]) + 1):
            width = max(
                (
                    len(str(row[column_index - 1] or ""))
                    for row in all_rows
                    if column_index <= len(row)
                ),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(width + 2, 12), 42
            )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def xlsx_from_csv(csv_text: str, *, sheet_name: str = "خروجی") -> bytes:
    """تبدیل خروجی CSV موجود به workbook، با حفظ متن‌ها و صفرهای ابتدایی شماره‌ها."""
    rows = list(csv.reader(io.StringIO(csv_text.lstrip("\ufeff"))))
    if not rows:
        return xlsx_bytes([], [], sheet_name=sheet_name)
    return xlsx_bytes(rows[0], rows[1:], sheet_name=sheet_name)
