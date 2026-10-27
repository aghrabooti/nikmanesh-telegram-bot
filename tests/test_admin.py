"""تست پنل مدیریت: ساخت/ویرایش/حذف، آمار، خروجی‌ها و بازه‌ی تاریخ."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook
from telegram import InputFile

from bot.config import content
from bot.services import seminars
from bot.services.database import get_db
from tests.helpers import (
    NON_ADMIN,
    callbacks,
    last_answer,
    last_document,
    last_edit,
    last_send,
    send_command,
    send_text,
    share_contact,
    tap,
    texts,
)


async def _create_via_panel(app, bot, *, title="سمینار آمادگی تیزهوشان", message="-"):
    """ساخت سمینار از پنل (فرم ۲ مرحله‌ای: اسم + پیام آخر)."""
    await tap(app, bot, "adm:sem:new")
    await send_text(app, bot, title)
    await send_text(app, bot, message)


async def _register_user(app, bot, seminar_id=1, *, user=NON_ADMIN, complete=True):
    """ثبت‌نام یک کاربر (کامل یا ناتمام تا مرحله‌ی پایه)."""
    await tap(app, bot, f"sem:reg:{seminar_id}", user=user)
    await send_text(app, bot, "علی محمدی", user=user)
    if not complete:
        return
    await send_text(app, bot, "کنکوری ۴۰۶", user=user)
    await send_text(app, bot, "تجربی", user=user)
    await send_text(app, bot, "تهران", user=user)
    await share_contact(app, bot, user=user, phone="09121234567")
    await send_text(app, bot, "09123456789", user=user)


async def test_admin_overview_callbacks(bot_app):
    app, bot = bot_app
    await send_command(app, bot, "/admin")
    sent = last_send(bot)
    assert sent["text"].startswith(content.ADMIN_MENU_HEADER)
    assert callbacks(sent["reply_markup"]) == [
        "adm:sems",
        "adm:bcast",
        "adm:users:csv",
        "adm:users:xlsx",
        "adm:ads",
        "adm:rem",
        "adm:about",
        "main",
    ]


async def test_create_seminar_full_flow(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot, message="سلام {name} جان، خوش آمدی")

    sem = seminars.get_seminar(1)
    assert sem["title"] == "سمینار آمادگی تیزهوشان"
    assert sem["message_template"] == "سلام {name} جان، خوش آمدی"
    assert sem["is_open"] == 1
    # آخرین پیام: صفحه‌ی سمینار ساخته‌شده
    assert content.ADMIN_CREATED_TEXT.format(title="سمینار آمادگی تیزهوشان") in texts(bot)[-1]


async def test_create_seminar_default_message(bot_app):
    """«-» در مرحله‌ی پیام آخر = پیام پیش‌فرض سراسری."""
    app, bot = bot_app
    await _create_via_panel(app, bot, title="سمینار مشاوره")
    assert seminars.get_seminar(1)["message_template"] == ""


async def test_create_message_too_long(bot_app):
    """پیام آخر بیش از ۸۰۰ نویسه → خطا و تکرار پرامپت."""
    app, bot = bot_app
    await tap(app, bot, "adm:sem:new")
    await send_text(app, bot, "سمینار تست")
    await send_text(app, bot, "x" * 801)
    assert last_send(bot)["text"] == content.ADMIN_TEMPLATE_INVALID
    await send_text(app, bot, "پیام کوتاه")
    assert seminars.get_seminar(1)["message_template"] == "پیام کوتاه"


async def test_create_title_validation(bot_app):
    app, bot = bot_app
    await tap(app, bot, "adm:sem:new")
    await send_text(app, bot, "ab")  # خیلی کوتاه
    assert last_send(bot)["text"] == content.ADMIN_TITLE_INVALID
    await send_text(app, bot, "سمینار تست")
    assert "۲ از ۲" in texts(bot)[-1]


async def test_sems_list_after_create(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await tap(app, bot, "adm:sems")

    edited = last_edit(bot)
    # آمار جداگانه‌ی هر سمینار در لیست
    assert "🟢 <b>سمینار آمادگی تیزهوشان</b> — ثبت‌نام <b>۰</b> (کامل <b>۰</b>)" in edited["text"]
    assert callbacks(edited["reply_markup"]) == ["adm:sem:1", "adm:sem:new", "main"]


async def test_seminar_page_shows_fields_and_stats(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)

    await tap(app, bot, "adm:sem:1")
    edited = last_edit(bot)
    assert "سمینار آمادگی تیزهوشان" in edited["text"]
    assert "پیام آخر: پیش‌فرض" in edited["text"]
    assert "کامل: <b>۱</b>" in edited["text"]
    assert "نیمه‌کاره: <b>۰</b>" in edited["text"]
    assert "قیف:" in edited["text"]
    assert "شماره جایگزین" in edited["text"]  # برچسب مرحله‌ی آخر


async def test_edit_menu_fields(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await tap(app, bot, "adm:sem:1")
    await tap(app, bot, "adm:sem:e:1")

    edited = last_edit(bot)
    assert edited["text"].startswith(content.ADMIN_EDIT_HEADER.split("\n")[0])
    assert callbacks(edited["reply_markup"]) == [
        "adm:sem:ef:1:title",
        "adm:sem:ef:1:message_template",
        "adm:sem:ef:1:starts_at",
        "adm:sem:1",
        "main",
    ]


async def test_edit_title(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await tap(app, bot, "adm:sem:e:1")
    await tap(app, bot, "adm:sem:ef:1:title")

    assert last_edit(bot)["text"] == content.ADMIN_EDIT_PROMPTS["title"]
    await send_text(app, bot, "عنوان جدید")
    assert seminars.get_seminar(1)["title"] == "عنوان جدید"
    assert "ذخیره شد" in texts(bot)[-2]
    assert content.ADMIN_EDIT_HEADER.split("\n")[0] in texts(bot)[-1]


async def test_open_close_toggle(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await tap(app, bot, "adm:sem:1")
    await tap(app, bot, "adm:sem:open:1")

    assert seminars.get_seminar(1)["is_open"] == 0
    assert last_answer(bot)["text"] == content.ADMIN_CLOSED_DONE
    await tap(app, bot, "adm:sem:open:1")
    assert seminars.get_seminar(1)["is_open"] == 1
    assert last_answer(bot)["text"] == content.ADMIN_OPENED_DONE


async def test_delete_seminar_two_step(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await tap(app, bot, "adm:sem:1")
    await tap(app, bot, "adm:sem:del:1")

    edited = last_edit(bot)
    assert content.ADMIN_DELETE_HEADER.split("\n")[0] in edited["text"]
    assert "سمینار آمادگی تیزهوشان" in edited["text"]

    await tap(app, bot, "adm:sem:del2:1")
    assert seminars.get_seminar(1) is None
    assert last_edit(bot)["text"].startswith(content.ADMIN_SEMS_HEADER.split("\n")[0])


async def test_delete_removes_registrations(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)
    await tap(app, bot, "adm:sem:1")
    await tap(app, bot, "adm:sem:del:1")
    await tap(app, bot, "adm:sem:del2:1")
    assert seminars.list_registrations(1) == []


async def test_regs_page_counts(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)  # کامل
    await _create_via_panel(app, bot, title="سمینار دوم")
    await _register_user(app, bot, 2, complete=False)  # ناتمام

    await tap(app, bot, "adm:regs:1")
    edited = last_edit(bot)
    assert "کل: <b>۱</b>" in edited["text"]
    assert "کامل: <b>۱</b>" in edited["text"]
    assert callbacks(edited["reply_markup"])[:5] == [
        "adm:regs:msg:1",
        "adm:regs:csv:1",
        "adm:regs:csv-complete:1",
        "adm:regs:xlsx:1",
        "adm:regs:xlsx-complete:1",
    ]


async def test_regs_message_output(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)
    await tap(app, bot, "adm:regs:1")
    await tap(app, bot, "adm:regs:msg:1")

    out = last_send(bot)["text"]
    assert "علی محمدی" in out
    assert "کنکوری ۴۰۶" in out
    assert "تجربی" in out
    assert "تهران" in out
    assert "09123456789" in out
    assert "✅" in out
    assert "کل: ۱ | کامل: ۱ | نیمه‌کاره: ۰" in out


async def test_regs_csv_output(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)
    await tap(app, bot, "adm:regs:1")
    await tap(app, bot, "adm:regs:csv:1")

    doc = last_document(bot)
    assert doc["filename"] == "seminar-1-registrations-all.csv"
    upload = InputFile(doc["document"], filename=doc["filename"])
    assert upload.filename == doc["filename"] and upload.mimetype == "text/csv"
    raw = doc["document"].getvalue()
    text = raw.decode("utf-8")
    assert text.startswith("\ufeff")  # BOM برای Excel فارسی
    rows = list(csv.reader(io.StringIO(text.lstrip("\ufeff"))))
    assert rows[0] == [
        "ردیف",
        "نام",
        "پایه",
        "رشته",
        "شهر",
        "شماره تلگرام",
        "شماره جایگزین",
        "وضعیت",
        "تاریخ ثبت",
    ]
    assert rows[1][1] == "علی محمدی"
    assert rows[1][2] == "کنکوری ۴۰۶"
    assert rows[1][3] == "تجربی"
    assert rows[1][5] == "+989121234567"
    assert rows[1][6] == "09123456789"
    assert rows[1][7] == "کامل"


async def test_regs_csv_complete_only_excludes_partial(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    seminars.ensure_registration(1001, 1)
    seminars.save_registration(
        1001,
        1,
        next_step=seminars.STEP_ALTERNATE_PHONE,
        name="ثبت کامل",
        alternate_phone="09111111111",
    )
    seminars.complete_registration(1001, 1)
    seminars.ensure_registration(1002, 1)
    seminars.save_registration(1002, 1, next_step=seminars.STEP_NAME, name="ثبت ناقص")

    await tap(app, bot, "adm:regs:1")
    await tap(app, bot, "adm:regs:csv-complete:1")
    doc = last_document(bot)
    assert doc["filename"] == "seminar-1-registrations-completed.csv"
    rows = list(
        csv.reader(io.StringIO(doc["document"].getvalue().decode("utf-8").lstrip("\ufeff")))
    )
    assert len(rows) == 2  # سرصفحه + تنها ثبت کامل
    assert rows[1][1] == "ثبت کامل"
    assert rows[1][7] == "کامل"


async def test_started_users_csv(bot_app):
    app, bot = bot_app
    await send_command(app, bot, "/start", user=NON_ADMIN)
    await tap(app, bot, "adm:users:csv")

    doc = last_document(bot)
    assert doc["filename"] == "users-started.csv"
    upload = InputFile(doc["document"], filename=doc["filename"])
    assert upload.filename == doc["filename"] and upload.mimetype == "text/csv"
    rows = list(
        csv.reader(io.StringIO(doc["document"].getvalue().decode("utf-8").lstrip("\ufeff")))
    )
    assert rows[0] == ["شناسه تلگرام", "نام", "نام کاربری", "زمان شروع", "آخرین بازدید"]
    assert rows[1][0] == str(NON_ADMIN.id)
    assert rows[1][1] == NON_ADMIN.first_name


async def test_regs_xlsx_respects_range_and_completed_filter(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    seminars.ensure_registration(1001, 1)
    seminars.save_registration(
        1001,
        1,
        next_step=seminars.STEP_ALTERNATE_PHONE,
        name="ثبت امروز",
        alternate_phone="09111111111",
    )
    seminars.complete_registration(1001, 1)
    seminars.ensure_registration(1002, 1)
    seminars.save_registration(1002, 1, next_step=seminars.STEP_NAME, name="ثبت قدیمی")
    get_db().execute(
        "UPDATE registrations SET created_at = ?, updated_at = ? WHERE user_id = ?",
        ("2000-01-01 00:00:00", "2000-01-01 00:00:00", 1002),
    )
    get_db().commit()

    await tap(app, bot, "adm:regs:1")
    await tap(app, bot, "adm:range:today")
    await tap(app, bot, "adm:regs:xlsx:1")
    doc = last_document(bot)
    assert doc["filename"] == "seminar-1-registrations-all.xlsx"
    upload = InputFile(doc["document"], filename=doc["filename"])
    assert upload.filename == doc["filename"]
    assert upload.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(io.BytesIO(upload.input_file_content), data_only=True)
    sheet = workbook.active
    rows = list(sheet.values)
    assert sheet.sheet_view.rightToLeft is True
    assert rows[0][6] == "شماره جایگزین"
    assert len(rows) == 2 and rows[1][1] == "ثبت امروز"

    await tap(app, bot, "adm:regs:xlsx-complete:1")
    doc = last_document(bot)
    workbook = load_workbook(io.BytesIO(doc["document"].getvalue()), read_only=True, data_only=True)
    rows = list(workbook.active.values)
    assert len(rows) == 2 and rows[1][1] == "ثبت امروز" and rows[1][7] == "کامل"


async def test_started_users_xlsx(bot_app):
    app, bot = bot_app
    await send_command(app, bot, "/start", user=NON_ADMIN)
    await tap(app, bot, "adm:users:xlsx")

    doc = last_document(bot)
    assert doc["filename"] == "users-started.xlsx"
    workbook = load_workbook(io.BytesIO(doc["document"].getvalue()), read_only=True)
    rows = list(workbook.active.values)
    assert rows[0] == ("شناسه تلگرام", "نام", "نام کاربری", "زمان شروع", "آخرین بازدید")
    assert rows[1][0] == NON_ADMIN.id


async def test_regs_range_quick(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)
    await tap(app, bot, "adm:regs:1")

    await tap(app, bot, "adm:range:all")
    assert "همه‌ی زمان‌ها" in last_edit(bot)["text"]
    assert "کل: <b>۱</b>" in last_edit(bot)["text"]

    await tap(app, bot, "adm:range:today")
    assert "امروز" in last_edit(bot)["text"]
    assert "کل: <b>۱</b>" in last_edit(bot)["text"]

    await tap(app, bot, "adm:range:90")
    assert "۹۰ روز اخیر" in last_edit(bot)["text"]


async def test_regs_range_custom(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)
    await tap(app, bot, "adm:regs:1")
    await tap(app, bot, "adm:range:custom")

    assert "دو تاریخ شمسی" in last_edit(bot)["text"]
    await send_text(app, bot, "۱۴۰۴/۰۱/۰۱ تا ۱۴۰۵/۰۱/۰۱")
    assert "از ۱۴۰۴/۰۱/۰۱ تا ۱۴۰۵/۰۱/۰۱" in texts(bot)[-1]

    # فرمت غلط → خطا و تکرار پرامپت
    await tap(app, bot, "adm:range:custom")
    await send_text(app, bot, "تاریخ الکی")
    assert last_send(bot)["text"] == content.ADMIN_RANGE_INVALID


async def test_regs_message_respects_range(bot_app):
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)
    await tap(app, bot, "adm:regs:1")
    await tap(app, bot, "adm:range:custom")
    await send_text(app, bot, "۱۴۰۰/۰۱/۰۱ تا ۱۴۰۰/۰۱/۰۲")  # بازه‌ی گذشته
    await tap(app, bot, "adm:regs:msg:1")
    assert last_send(bot)["text"] == content.ADMIN_REGS_EMPTY


async def test_admin_fallback_no_state(bot_app):
    app, bot = bot_app
    await send_text(app, bot, "سلام")  # مدیر بدون state فعال
    assert last_send(bot)["text"] == content.FALLBACK_TEXT


async def test_non_admin_admin_callbacks_denied(bot_app):
    app, bot = bot_app
    await tap(app, bot, "adm:sems", user=NON_ADMIN)
    assert last_answer(bot)["text"] == content.ADMIN_ONLY_TEXT
    assert bot.edit_message_text.await_count == 0
    await tap(app, bot, "adm:sem:new", user=NON_ADMIN)
    assert seminars.list_seminars() == []


async def test_admin_flow_back_and_main_menu(bot_app):
    """ناوبری پنل: بازگشت و منوی اصلی."""
    app, bot = bot_app
    await _create_via_panel(app, bot)
    await _register_user(app, bot)
    await tap(app, bot, "adm:sem:1")
    await tap(app, bot, "adm:sem:e:1")
    await tap(app, bot, "adm:sem:1")  # 🔙 از منوی ویرایش
    assert "قیف:" in last_edit(bot)["text"]
    await tap(app, bot, "adm:sems")  # 🔙 از صفحه‌ی سمینار
    assert "مدیریت سمینارها" in last_edit(bot)["text"]
    await tap(app, bot, "main")  # 🏠 منوی اصلی کاربر
    assert last_edit(bot)["text"] == content.WELCOME_TEXT
