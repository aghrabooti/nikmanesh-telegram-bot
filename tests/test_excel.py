"""تست‌های ساخت فایل Excel واقعی و امن برای متن‌های کاربر."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from bot.utils.excel import xlsx_bytes, xlsx_from_csv


def test_xlsx_bytes_creates_right_to_left_workbook_and_keeps_text_safe():
    payload = xlsx_bytes(
        ["نام", "یادداشت"],
        [["آزمایشی", '=HYPERLINK("https://example.com","click")']],
        sheet_name="گزارش",
    )

    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    sheet = workbook.active
    assert sheet.title == "گزارش"
    assert sheet.sheet_view.rightToLeft is True
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B2"
    assert sheet.cell(1, 1).value == "نام"
    assert sheet.cell(2, 2).value.startswith("=HYPERLINK")
    assert sheet.cell(2, 2).data_type == "s"


def test_xlsx_from_csv_parses_farsi_headers_and_rows():
    payload = xlsx_from_csv("نام,شماره جایگزین\nسارا,09123456789\n", sheet_name="ثبت‌نام‌ها")

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    rows = list(workbook.active.values)
    assert rows == [("نام", "شماره جایگزین"), ("سارا", "09123456789")]
