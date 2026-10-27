"""تست لینک‌های رفرال اسم‌دار — پورت سیستم تیزگام (ref_<code>)."""

from __future__ import annotations

import csv
import io
import sqlite3

from openpyxl import load_workbook
from telegram import InputFile

from bot.config import content
from bot.constants import REF_START_PREFIX
from bot.services import referrals, seminars
from bot.services.database import close_db, init_db
from tests.helpers import (
    NON_ADMIN,
    USER,
    callbacks,
    labels,
    last_document,
    last_edit,
    last_send,
    send_command,
    send_text,
    share_contact,
    tap,
)


def _create(**kwargs):
    defaults = dict(title="وبینار تست", link="https://t.me/room", levels=[])
    defaults.update(kwargs)
    return seminars.create_seminar(**defaults)


async def _register(app, bot, seminar_id, *, user=NON_ADMIN, complete=True, name="سارا احمدی"):
    """ثبت‌نام کامل یا ناتمام (فقط تا مرحله‌ی نام)."""
    await tap(app, bot, f"sem:reg:{seminar_id}", user=user)
    await send_text(app, bot, name, user=user)
    if not complete:
        return
    await send_text(app, bot, "کنکوری ۴۰۷", user=user)
    await send_text(app, bot, "ریاضی", user=user)
    await send_text(app, bot, "شیراز", user=user)
    await share_contact(app, bot, user=user, phone="09121234567")
    await send_text(app, bot, "۰۹۱۲۳۴۵۶۷۸۹", user=user)


# ── ساخت لینک: چندتا، با اسم دلخواه ──────────────────────────────────────


def test_create_multiple_named_links(bot_app):
    app, bot = bot_app
    a = referrals.create_link("کانال ۱ — ریاضی")
    b = referrals.create_link("استوری اینستاگرام")
    c = referrals.create_link("گروه داوطلبان")
    assert a["label"] == "کانال ۱ — ریاضی"
    assert b["label"] == "استوری اینستاگرام"
    assert c["label"] == "گروه داوطلبان"
    codes = {a["code"], b["code"], c["code"]}
    assert len(codes) == 3 and all(len(x) == 6 for x in codes)
    assert all(lnk["is_active"] for lnk in (a, b, c))
    # لیست: جدیدترین اول
    assert [lnk["link_id"] for lnk in referrals.links()] == [
        c["link_id"],
        b["link_id"],
        a["link_id"],
    ]


def test_link_url_format(bot_app):
    app, bot = bot_app
    link = referrals.create_link("تست")
    assert referrals.link_url("nikmanesh_bot", link["code"]) == (
        f"https://t.me/nikmanesh_bot?start={REF_START_PREFIX}{link['code']}"
    )


# ── اتریبیوشن first-touch ────────────────────────────────────────────────


async def test_start_with_ref_attributes_and_welcome_normal(bot_app):
    app, bot = bot_app
    s = _create()
    link = referrals.create_link("کانال تست")
    await send_command(app, bot, f"/start {REF_START_PREFIX}{link['code']}", user=NON_ADMIN)
    # استارت ثبت شده؛ خوش‌آمد عادی آمده (دکمه‌ی شیشه‌ای = وبینار فعال سراسری)
    assert referrals.funnel(link["link_id"])["starts"] == 1
    welcome = bot.send_message.await_args_list[-2].kwargs
    assert callbacks(welcome["reply_markup"]) == [f"sem:v:{s['id']}"]


async def test_first_touch_only_first_link_counts(bot_app):
    app, bot = bot_app
    first = referrals.create_link("لینک اول")
    second = referrals.create_link("لینک دوم")
    await send_command(app, bot, f"/start {REF_START_PREFIX}{first['code']}", user=NON_ADMIN)
    await send_command(app, bot, f"/start {REF_START_PREFIX}{second['code']}", user=NON_ADMIN)
    await send_command(app, bot, "/start", user=NON_ADMIN)
    assert referrals.funnel(first["link_id"])["starts"] == 1
    assert referrals.funnel(second["link_id"])["starts"] == 0


async def test_invalid_or_inactive_link_not_counted(bot_app):
    app, bot = bot_app
    link = referrals.create_link("قابل خاموش")
    referrals.set_active(link["link_id"], False)
    await send_command(app, bot, f"/start {REF_START_PREFIX}{link['code']}", user=NON_ADMIN)
    await send_command(app, bot, f"/start {REF_START_PREFIX}zzzzzz", user=NON_ADMIN)
    assert referrals.funnel(link["link_id"])["starts"] == 0


# ── قیف هر لینک + ارگانیک ────────────────────────────────────────────────


async def test_funnel_per_link_and_organic(bot_app):
    app, bot = bot_app
    s = _create()
    via_link = referrals.create_link("کمپین اصلی")
    # کاربر ۱: از لینک، ثبت‌نام کامل
    await send_command(app, bot, f"/start {REF_START_PREFIX}{via_link['code']}", user=NON_ADMIN)
    await _register(app, bot, s["id"], user=NON_ADMIN)
    # کاربر ۲: از لینک، نیمه‌کاره (تا مرحله‌ی نام)
    await send_command(app, bot, f"/start {REF_START_PREFIX}{via_link['code']}", user=USER)
    await _register(app, bot, s["id"], user=USER, complete=False, name="مهمان تست")
    # ثبت‌نام ارگانیک: مستقیم از سرویس، بدون /start با لینک
    third = referrals.create_link("بی‌استفاده")
    seminars.ensure_registration(999001, s["id"])
    seminars.save_registration(999001, s["id"], next_step=seminars.STEP_GRADE, name="قدیمی")

    f = referrals.funnel(via_link["link_id"])
    assert f["starts"] == 2
    assert f["steps"][seminars.STEP_NAME] == 2
    assert f["steps"][seminars.STEP_GRADE] == 2  # کامل‌شده + نیمه‌کاره‌ی در مرحله‌ی پایه
    assert f["steps"][seminars.STEP_FIELD] == 1  # فقط کامل‌شده گذشته
    assert f["completed"] == 1
    assert f["grades"]["کنکوری ۴۰۷"] == 1

    organic = referrals.funnel(None)
    assert organic["starts"] == 1
    assert organic["completed"] == 0

    unused = referrals.funnel(third["link_id"])
    assert unused["starts"] == 0 and unused["completed"] == 0


# ── پنل ادمین: هاب، ساخت، جزئیات، خاموش/روشن، CSV ────────────────────────


async def test_admin_ads_hub_lists_named_links_and_new_button(bot_app):
    app, bot = bot_app
    a = referrals.create_link("کانال فیزیک")
    b = referrals.create_link("استوری")
    referrals.set_active(b["link_id"], False)
    await tap(app, bot, "adm:ads", user=USER)
    sent = last_edit(bot)
    assert content.ADS_HUB_TITLE in sent["text"]
    lbls = labels(sent["reply_markup"])
    assert f"کانال فیزیک — {a['code']}" in lbls
    assert f"⏸ استوری — {b['code']}" in lbls  # غیرفعال = پیشوند ⏸
    assert "adm:ads:new" in callbacks(sent["reply_markup"])
    assert "adm:ads:msg" in callbacks(sent["reply_markup"])
    assert "adm:ads:csv" in callbacks(sent["reply_markup"])
    assert "adm:ads:xlsx" in callbacks(sent["reply_markup"])


async def test_admin_create_link_with_label_flow(bot_app):
    """ساخت لینک از پنل: دکمه‌ی جدید → ارسال اسم → لینک کامل + صفحه‌ی جزئیات."""
    app, bot = bot_app
    bot.username = "nikmanesh_bot"
    _create()
    await tap(app, bot, "adm:ads", user=USER)
    await tap(app, bot, "adm:ads:new", user=USER)
    assert content.ADS_NEW_PROMPT in last_edit(bot)["text"]

    await send_text(app, bot, "کانال کنکور ریاضی", user=USER)
    link = referrals.links()[0]
    assert link["label"] == "کانال کنکور ریاضی"
    created = last_send(bot)
    # پیام تأیید شامل لینک کامل و رمز است.
    assert "کانال کنکور ریاضی" in created["text"]
    assert f"https://t.me/nikmanesh_bot?start={REF_START_PREFIX}" in created["text"]
    # همان پیام پرامپت به هاب به‌روز‌شده برمی‌گردد و لینک تازه دکمه دارد.
    updated_hub = last_edit(bot)
    assert content.ADS_HUB_TITLE in updated_hub["text"]
    assert f"adm:ads:v:{link['link_id']}" in callbacks(updated_hub["reply_markup"])


async def test_admin_link_detail_and_toggle(bot_app):
    app, bot = bot_app
    bot.username = "nikmanesh_bot"
    s = _create()
    link = referrals.create_link("کمپین اینستا")
    await send_command(app, bot, f"/start {REF_START_PREFIX}{link['code']}", user=NON_ADMIN)
    await _register(app, bot, s["id"], user=NON_ADMIN)

    await tap(app, bot, f"adm:ads:v:{link['link_id']}", user=USER)
    detail = last_edit(bot)
    assert "کمپین اینستا" in detail["text"]
    assert referrals.link_url("nikmanesh_bot", link["code"]) in detail["text"]
    assert "قیف ثبت‌نام" in detail["text"] and "✅ ۱" in detail["text"]

    # خاموش کردن → استارت بعدی با آن شمرده نمی‌شود
    await tap(app, bot, f"adm:ads:t:{link['link_id']}", user=USER)
    assert "غیرفعال" in last_edit(bot)["text"]
    assert referrals.links()[0]["is_active"] == 0
    await send_command(app, bot, f"/start {REF_START_PREFIX}{link['code']}", user=USER)
    assert referrals.funnel(link["link_id"])["starts"] == 1  # فقط استارت اول


async def test_admin_link_regs_message_and_csv(bot_app):
    """صفحه‌ی لینک: لیست ثبت‌نامی‌های همان لینک (اسم+شماره‌ها) + CSV جدا."""
    app, bot = bot_app
    s = _create()
    link = referrals.create_link("کمپین الف")
    other = referrals.create_link("کمپین ب")
    # کاربر ۱: از لینک الف، کامل
    await send_command(app, bot, f"/start {REF_START_PREFIX}{link['code']}", user=NON_ADMIN)
    await _register(app, bot, s["id"], user=NON_ADMIN)
    # کاربر ۲: از لینک ب، نیمه‌کاره (نباید در لیست لینک الف باشد)
    await send_command(app, bot, f"/start {REF_START_PREFIX}{other['code']}", user=USER)
    await _register(app, bot, s["id"], user=USER, complete=False, name="مهمان نیمه")

    # دکمه‌ها در صفحه‌ی لینک
    await tap(app, bot, f"adm:ads:v:{link['link_id']}", user=USER)
    cbs = callbacks(last_edit(bot)["reply_markup"])
    assert f"adm:ads:rg:{link['link_id']}" in cbs
    assert f"adm:ads:rgc:{link['link_id']}" in cbs
    assert f"adm:ads:rgx:{link['link_id']}" in cbs

    # پیام لیست: فقط ثبت‌نامیِ همین لینک، با اسم و شماره‌ها
    await tap(app, bot, f"adm:ads:rg:{link['link_id']}", user=USER)
    text = last_send(bot)["text"]
    assert "کمپین الف" in text
    assert "سارا احمدی" in text and "مهمان نیمه" not in text
    assert "+989121234567" in text and "09123456789" in text
    assert "✅" in text

    # CSV همین لینک
    await tap(app, bot, f"adm:ads:rgc:{link['link_id']}", user=USER)
    doc = last_document(bot)
    assert doc["filename"] == f"ref-{link['code']}-regs.csv"
    rows = list(
        csv.reader(io.StringIO(doc["document"].getvalue().decode("utf-8").lstrip("\ufeff")))
    )
    assert rows[0][1:7] == ["نام", "پایه", "رشته", "شهر", "شماره تلگرام", "شماره جایگزین"]
    assert len(rows) == 2  # سرصفحه + فقط کاربر ۱
    assert rows[1][1] == "سارا احمدی"
    assert rows[1][5] == "+989121234567" and rows[1][6] == "09123456789"
    assert rows[1][7] == "کامل"

    # Excel واقعی همین لینک
    await tap(app, bot, f"adm:ads:rgx:{link['link_id']}", user=USER)
    doc = last_document(bot)
    assert doc["filename"] == f"ref-{link['code']}-regs.xlsx"
    upload = InputFile(doc["document"], filename=doc["filename"])
    assert upload.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(io.BytesIO(upload.input_file_content), read_only=True, data_only=True)
    excel_rows = list(workbook.active.values)
    assert excel_rows[0][6] == "شماره جایگزین"
    assert excel_rows[1][1] == "سارا احمدی" and excel_rows[1][6] == "09123456789"


async def test_admin_ads_msg_full_stats(bot_app):
    app, bot = bot_app
    _create()
    link = referrals.create_link("لینک الف")
    await tap(app, bot, "adm:ads:msg", user=USER)
    text = last_send(bot)["text"]
    assert content.ADS_STATS_TITLE in text
    assert content.ADS_ORGANIC_TITLE in text  # ارگانیک هم هست
    assert "لینک الف" in text and f"ref_{link['code']}" in text


async def test_admin_ads_csv_document(bot_app):
    app, bot = bot_app
    bot.username = "nikmanesh_bot"
    referrals.create_link("کانال الف")
    referrals.create_link("کانال ب")
    await tap(app, bot, "adm:ads:csv", user=USER)
    doc = last_document(bot)
    assert doc["filename"] == "nikmanesh-ads.csv"
    raw = doc["document"].getvalue().decode("utf-8")
    rows = list(csv.reader(io.StringIO(raw.lstrip("\ufeff"))))
    assert rows[0][:6] == ["code", "label", "active", "created_at", "link", "starts"]
    assert [r[1] for r in rows[1:]] == ["کانال ب", "کانال الف"]  # جدیدترین اول
    assert all(r[4].startswith("https://t.me/nikmanesh_bot?start=ref_") for r in rows[1:])

    await tap(app, bot, "adm:ads:xlsx", user=USER)
    doc = last_document(bot)
    assert doc["filename"] == "nikmanesh-ads.xlsx"
    workbook = load_workbook(io.BytesIO(doc["document"].getvalue()), read_only=True)
    excel_rows = list(workbook.active.values)
    assert excel_rows[0][:6] == ("code", "label", "active", "created_at", "link", "starts")
    assert [row[1] for row in excel_rows[1:]] == ["کانال ب", "کانال الف"]


# ── مهاجرت دیتابیس قدیمی ────────────────────────────────────────────────


def test_legacy_db_gets_link_id_column(tmp_path):
    """registrations قدیمی بدون link_id → ALTER خودکار؛ اتریبیوشن کار می‌کند."""
    db_path = tmp_path / "legacy.db"
    conn = sqlite3.connect(db_path)
    conn.execute(
        "CREATE TABLE registrations ("
        "user_id INTEGER NOT NULL, seminar_id INTEGER NOT NULL, name TEXT, "
        "current_step TEXT, is_complete INTEGER, created_at TEXT, updated_at TEXT, "
        "UNIQUE (user_id, seminar_id))"
    )
    conn.commit()
    conn.close()

    init_db(db_path)
    seminars.create_seminar(title="سمینار قدیمی")
    link = referrals.create_link("لینک پس از مهاجرت")
    referrals.get_db().execute(
        "INSERT OR IGNORE INTO referral_starts (user_id, link_id, started_at) VALUES (?, ?, ?)",
        (555, link["link_id"], "2026-01-01 10:00:00"),
    )
    referrals.get_db().commit()
    seminars.ensure_registration(555, 1)
    reg = seminars.get_registration(555, 1)
    assert reg["link_id"] == link["link_id"]
    close_db()
